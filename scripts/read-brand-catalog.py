"""Read-only extraction: no costs, prices or stock leave the source workbook."""
import hashlib
import json
import pathlib
import sys
from openpyxl import load_workbook

source = pathlib.Path(sys.argv[1])
workbook = load_workbook(source, read_only=True, data_only=True)
sheet = workbook['Hoja1']
rows = sheet.iter_rows(values_only=True)
headers = [str(v or '').strip() for v in next(rows)]
required = ['Referencia', 'Descripción', 'LINEA', 'CATEGORIA', 'MARCA']
if any(name not in headers for name in required):
    raise ValueError('El Excel no tiene las columnas esperadas del catálogo.')
result, references = [], set()
for number, row in enumerate(rows, 2):
    if not any(v is not None for v in row):
        continue
    fields = {name: str(row[headers.index(name)] or '').strip() for name in required}
    if any(not fields[name] for name in ['Referencia', 'Descripción', 'MARCA']):
        raise ValueError(f'Faltan referencia, descripción o marca en fila {number}.')
    if fields['Referencia'].upper() in references:
        raise ValueError(f'Referencia repetida en fila {number}.')
    references.add(fields['Referencia'].upper())
    for name, limit in [('Referencia', 100), ('Descripción', 300), ('LINEA', 100), ('CATEGORIA', 100), ('MARCA', 100)]:
        if len(fields[name]) > limit:
            raise ValueError(f'{name} excede la longitud permitida en fila {number}.')
    held = fields['MARCA'].upper() in {'HAEB', 'HCEB'}
    result.append(dict(referencia=fields['Referencia'], descripcion=fields['Descripción'],
                       marca=fields['MARCA'], linea=fields['LINEA'], categoria=fields['CATEGORIA'],
                       habilitada=not held, motivo='Marca pendiente de confirmar en el archivo original' if held else None,
                       fila=number))
workbook.close()
if not result:
    raise ValueError('El catálogo está vacío.')
print(json.dumps(dict(archivo=source.name, sha256=hashlib.sha256(source.read_bytes()).hexdigest(), filas=result), ensure_ascii=True))
