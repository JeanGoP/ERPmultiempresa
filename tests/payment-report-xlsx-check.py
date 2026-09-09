"""Inspección independiente de los XLSX descargados por la prueba de navegador."""
from pathlib import Path
import zipfile
import openpyxl

root = Path(__file__).resolve().parent.parent / 'tmp' / 'payment-report-qa'
with zipfile.ZipFile(root / 'relacion.xlsx') as archive:
    assert archive.testzip() is None, 'CRC o estructura ZIP incorrecta'
book = openpyxl.load_workbook(root / 'relacion.xlsx')
values = openpyxl.load_workbook(root / 'relacion.xlsx', data_only=True)
assert book.sheetnames == ['Facturas', 'Pagos y notas']
invoices = book['Facturas']
assert invoices['B12'].value == '001234' and invoices['B12'].data_type == 's'
assert invoices['B12'].number_format == '@'
assert invoices['C12'].value == 1000000 and invoices['C12'].data_type == 'n'
assert invoices['E12'].value == '=ROUND(C12-D12,2)'
assert values['Facturas']['K13'].value == 746200
assert invoices.freeze_panes == 'A12' and invoices.auto_filter.ref == 'B11:K12'
assert not invoices.sheet_view.showGridLines
assert invoices['C12'].number_format.startswith('#,##0.00')
payments = book['Pagos y notas']
assert payments['B12'].value.year == 2026 and payments['B12'].number_format == 'dd/mm/yyyy'
assert payments['C12'].value == '001234' and payments['C12'].data_type == 's'
assert payments['F53'].data_type == 'f' and values['Pagos y notas']['G53'].value == 420000
assert 'NC-1' in [payments.cell(r, 5).value for r in range(12, 53)]
assert payments['B12'].fill.fgColor.rgb != payments['B13'].fill.fgColor.rgb
hostile = openpyxl.load_workbook(root / 'texto-seguro.xlsx')
assert hostile['Facturas']['B12'].data_type == 's'
assert hostile['Facturas']['B12'].value.startswith('=HYPERLINK')
assert 'Nombre & <etiqueta> "literal"' in hostile['Facturas']['B4'].value
empty = openpyxl.load_workbook(root / 'sin-pagos.xlsx')
assert 'Sin pagos' in empty['Pagos y notas']['B12'].value
print('XLSX: ZIP, hojas, tipos, fórmulas, totales, formato, fechas, ceros iniciales y textos seguros correctos.')
